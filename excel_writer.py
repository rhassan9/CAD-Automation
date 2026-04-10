import openpyxl
import os
import shutil

class ExcelWriter:
    def __init__(self, output_path, template_path=None):
        self.output_path = output_path
        
        # Template Fallback Generation
        if template_path and os.path.exists(template_path):
            shutil.copy(template_path, output_path)
            self.wb = openpyxl.load_workbook(output_path)
            print(f"Loaded existing template: {output_path}")
        else:
            self.wb = openpyxl.Workbook()
            sheets = ['House Count', 'CABLE TO 1X8 SPLICING', 'SPLICING 1X8 TO 1X4 SPLITS']
            for name in sheets:
                self.wb.create_sheet(title=name)
            if 'Sheet' in self.wb.sheetnames:
                self.wb.remove(self.wb['Sheet'])
            print(f"Generated Raw Fallback Template: {output_path}")
            
    def populate_house_count(self, house_data):
        print("Populating House Count sheet...")
        if 'House Count' not in self.wb.sheetnames: return
        sheet = self.wb['House Count']
        
        from openpyxl.styles import Font, Alignment, Border, Side
        
        calibri_11 = Font(name='Calibri', size=11)
        center_align = Alignment(horizontal='center', vertical='center')
        
        # House Count Formatting Logic
        thin_side = Side(style='thin', color='000000')
        thick_side = Side(style='medium', color='000000')
        
        standard_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
        thick_bottom_border = Border(top=thin_side, bottom=thick_side, left=thin_side, right=thin_side)
        
        # Deduplicate houses based on ADDRESS + UNIT
        unique_houses = {}
        for data in house_data:
            key = f"{data['ADDRESS']}||{data.get('UNIT', '')}"
            if key not in unique_houses:
                unique_houses[key] = data
                
        # Sort alphabetically by address
        sorted_houses = sorted(unique_houses.values(), key=lambda x: x['ADDRESS'])
        
        row = 2
        for idx, data in enumerate(sorted_houses):
            # Check if this row is a multiple of 12 (0-indexed idx)
            is_12th_row = ((idx + 1) % 12 == 0)
            target_border = thick_bottom_border if is_12th_row else standard_border
            
            sheet.cell(row=row, column=1).value = data['ADDRESS']
            sheet.cell(row=row, column=2).value = "FUQUAY-VARINA"
            sheet.cell(row=row, column=3).value = "NORTH CAROLINA"
            sheet.cell(row=row, column=4).value = "27526"
            sheet.cell(row=row, column=5).value = data.get('UNIT', '')
            sheet.cell(row=row, column=6).value = "RES-SFU"
            sheet.cell(row=row, column=7).value = "BURIED"
            sheet.cell(row=row, column=8).value = "Y"
            sheet.cell(row=row, column=9).value = ""
            sheet.cell(row=row, column=10).value = 1
            sheet.cell(row=row, column=11).value = 0
            
            # Apply formatting
            for col in range(1, 12):
                cell = sheet.cell(row=row, column=col)
                cell.font = calibri_11
                cell.alignment = center_align
                cell.border = target_border
                
            row += 1
            
    def populate_splices(self, splitters_1x8):
        print("Populating 1x8 Splice sheet dynamically...")
        sheet_1x8 = self.wb['CABLE TO 1X8 SPLICING']
        
        # We will hold back the 1x4 injection as requested until the next phase,
        # but we must save the extracted template formatting (Colors).
        self.extracted_1x8_template_data = {}

        # 1. Sort the DXF splitters numerically by placement_no
        sorted_splitters = []
        for b_name, struct in splitters_1x8.items():
            if struct.get('placement_no'):
                try:
                    num = int(struct['placement_no'])
                except ValueError:
                    num = 99999
                sorted_splitters.append((num, b_name, struct))
        sorted_splitters.sort(key=lambda x: x[0])

        for (num, base_name, struct) in sorted_splitters:
            placement_no = struct.get('placement_no', '')
            
            # --- 1X8 SPLICING BLOCK FINDER ---
            target_1x8_row = None
            for r in range(5, 500):
                # Search Column A for the index
                cell_val = str(sheet_1x8.cell(row=r, column=1).value).strip() if sheet_1x8.cell(row=r, column=1).value else None
                if cell_val == str(placement_no):
                    target_1x8_row = r
                    break
                    
            if target_1x8_row:
                # Capture Colors from Col B and C using copy module
                import copy
                color_1_fill = copy.copy(sheet_1x8.cell(row=target_1x8_row, column=2).fill)
                color_2_fill = copy.copy(sheet_1x8.cell(row=target_1x8_row, column=3).fill)
                
                # Store in class variable for later 1x4 Phase
                self.extracted_1x8_template_data[placement_no] = {
                    'color_1': color_1_fill,
                    'color_2': color_2_fill,
                    'row': target_1x8_row
                }

                # Write logic mapped specifically to the documentation
                # Column D: Extracted #
                sheet_1x8.cell(row=target_1x8_row, column=4).value = int(placement_no)
                
                # Column E: Address
                sheet_1x8.cell(row=target_1x8_row, column=5).value = struct.get('placement_address', '')
                
                # Column F: 1x8 Name
                sheet_1x8.cell(row=target_1x8_row, column=6).value = base_name

                # Injecting 1x8 Ports (G-N)
                for port in range(1, 9):
                    out_1x8_str = "SPARE"
                    if port in struct['ports']:
                        out_1x8_str = f"{base_name}, {port} ({base_name}-{port})"
                    
                    # Columns G is 7, H is 8, etc. -> 6 + port
                    sheet_1x8.cell(row=target_1x8_row, column=6+port).value = out_1x8_str
            else:
                print(f"WARNING: Placement {placement_no} not found in template Column A!")

    def populate_1x4_splits(self, splitters_1x8):
        print("Populating 1x8 to 1x4 Splice sheet dynamically...")
        if 'SPLICING 1X8 TO 1X4 SPLITS' not in self.wb.sheetnames:
            return
            
        sheet_1x4 = self.wb['SPLICING 1X8 TO 1X4 SPLITS']
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        
        header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        header_font = Font(name='Calibri', size=11, bold=True)
        center_align = Alignment(horizontal='center', vertical='center')
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        # 1. Sort the DXF splitters numerically by placement_no
        sorted_splitters = []
        for b_name, struct in splitters_1x8.items():
            if struct.get('placement_no'):
                try:
                    num = int(struct['placement_no'])
                except ValueError:
                    num = 99999
                sorted_splitters.append((num, b_name, struct))
        sorted_splitters.sort(key=lambda x: x[0])
        
        search_start_row = 1
        max_processed_row = 1
        
        for (num, base_name, struct) in sorted_splitters:
            placement_no = struct.get('placement_no', '')
            
            # --- ROBUST ANCHOR ROW FINDER ---
            anchor_row = None
            for r in range(search_start_row, 5000):
                cell_val = str(sheet_1x4.cell(row=r, column=1).value).strip() if sheet_1x4.cell(row=r, column=1).value else ''
                if cell_val == '1X8 SPLITTER PLACEMENT':
                    anchor_row = r + 1
                    search_start_row = anchor_row + 1 # Next search starts below this point
                    break
                    
            if not anchor_row:
                print(f"Info: Anchor row block for placement {placement_no} not found. Building fallback section...")
                # Find absolute bottom of existing template blocks to append to
                bottom_header = 1
                for r in range(1, 4000):
                    cell_val = str(sheet_1x4.cell(row=r, column=1).value).strip() if sheet_1x4.cell(row=r, column=1).value else ''
                    if cell_val == '1X8 SPLITTER PLACEMENT':
                        bottom_header = r
                        
                # New block goes below either the last template block or the last dynamically created block
                fallback_start = max(max_processed_row, bottom_header + 12) + 2 
                header_strip_row = fallback_start
                
                # --- APPLY FALLBACK FORMATTING ---
                # A to F merge for Main Header
                sheet_1x4.merge_cells(start_row=header_strip_row, start_column=1, end_row=header_strip_row, end_column=6)
                h_cell = sheet_1x4.cell(row=header_strip_row, column=1)
                h_cell.value = '1X8 SPLITTER PLACEMENT'
                h_cell.fill = header_fill
                h_cell.font = header_font
                h_cell.alignment = center_align
                
                for c in range(1, 9):
                    sheet_1x4.cell(row=header_strip_row, column=c).border = thin_border
                    
                anchor_row = header_strip_row + 1
                
                # Apply borders to Anchor Row
                for c in range(1, 9):
                    sheet_1x4.cell(row=anchor_row, column=c).border = thin_border
                    
                # Sub-headers (Merged Rows anchor+1 to anchor+2)
                sub_headers = ['SPLITTER NAME', 'T COUNT', '1X4 PLACEMENT ADDRESS', 'SPLIT', 'SERVICING', 'SERVICING', 'SERVICING', 'SERVICING']
                for c in range(1, 9):
                    sheet_1x4.merge_cells(start_row=anchor_row+1, start_column=c, end_row=anchor_row+2, end_column=c)
                    sub = sheet_1x4.cell(row=anchor_row+1, column=c)
                    sub.value = sub_headers[c-1]
                    sub.font = header_font
                    sub.alignment = center_align
                    sheet_1x4.cell(row=anchor_row+1, column=c).border = thin_border
                    sheet_1x4.cell(row=anchor_row+2, column=c).border = thin_border
                    
                # CRITICAL FIX: Ensure the next placement doesn't hijack the block we just built!
                search_start_row = anchor_row + 12
            
            # Record our progress to avoid overlaps
            max_processed_row = max(max_processed_row, anchor_row + 10)
            
            # --- START DATA INJECTION ---
            # Overwrite the user's nasty formula with our hardcoded, correct placement number
            p_cell = sheet_1x4.cell(row=anchor_row, column=1)
            p_cell.value = int(placement_no)
            p_cell.alignment = center_align
            
            # Retrieve colors saved during 1x8 processing
            if hasattr(self, 'extracted_1x8_template_data') and placement_no in self.extracted_1x8_template_data:
                import copy
                color_1 = self.extracted_1x8_template_data[placement_no].get('color_1')
                color_2 = self.extracted_1x8_template_data[placement_no].get('color_2')
                
                if color_1:
                    sheet_1x4.cell(row=anchor_row, column=3).fill = copy.copy(color_1)
                if color_2:
                    sheet_1x4.cell(row=anchor_row, column=4).fill = copy.copy(color_2)
                    
            # Column F: Parent Address
            addr_cell = sheet_1x4.cell(row=anchor_row, column=6)
            addr_cell.value = struct.get('placement_address', '')
            addr_cell.font = Font(name='Calibri', size=11, bold=True)
            addr_cell.alignment = center_align
            
            # --- ROW MAPPING (8 ROWS) ---
            t_counts = ["1-4", "5-8", "9-12", "13-16", "17-20", "21-24", "25-28", "29-32"]
            for port in range(1, 9):
                row_offset = anchor_row + 2 + port
                
                # Apply Borders to every cell in the data row globally
                for c in range(1, 9):
                    sheet_1x4.cell(row=row_offset, column=c).border = thin_border
                    sheet_1x4.cell(row=row_offset, column=c).alignment = center_align
                
                # Column B: T COUNT
                sheet_1x4.cell(row=row_offset, column=2).value = t_counts[port - 1]
                
                if port in struct['ports']:
                    port_data = struct['ports'][port]
                    
                    raw_name = port_data.get('1x4_name', '')
                    clean_name = raw_name.replace('1X4 SPLITTER ', '').strip()
                    sheet_1x4.cell(row=row_offset, column=1).value = clean_name
                    sheet_1x4.cell(row=row_offset, column=3).value = port_data.get('placement_address_1x4', '')
                    sheet_1x4.cell(row=row_offset, column=4).value = port_data.get('split_reference', '')
                    
                    houses = port_data.get('houses', [])
                    for i in range(4):
                        if i < len(houses) and houses[i]:
                            sheet_1x4.cell(row=row_offset, column=5 + i).value = houses[i]
                else:
                    sheet_1x4.cell(row=row_offset, column=4).value = 'SPARE'
                    # Strip any leftover template dummy data
                    sheet_1x4.cell(row=row_offset, column=1).value = None
                    sheet_1x4.cell(row=row_offset, column=3).value = None
                    sheet_1x4.cell(row=row_offset, column=5).value = None
                    sheet_1x4.cell(row=row_offset, column=6).value = None
                    sheet_1x4.cell(row=row_offset, column=7).value = None
                    sheet_1x4.cell(row=row_offset, column=8).value = None

    def populate_cable_sheet(self, data):
        print("Populating Cable Sheet dynamically...")
        if 'CABLE SHEET' not in self.wb.sheetnames:
            print("WARNING: CABLE SHEET not found in workbook.")
            return
            
        sheet = self.wb['CABLE SHEET']
        from openpyxl.styles import Alignment
        import math
        center_align = Alignment(horizontal='center', vertical='center')
        
        cable_spans = data.get('cable_spans', {})
        # Sort segments alphabetically represented by integer
        sorted_segs = sorted(cable_spans.keys())
        
        # Build flat list of all Splitter coordinates
        splitters = []
        for v in data.get('splitters_1x8', {}).values():
            splitters.append((v.get('x', 0), v.get('y', 0)))
        for v in data.get('splitters_1x4', []):
            splitters.append((v.get('X', 0), v.get('Y', 0)))
        
        for seg_idx, seg_num in enumerate(sorted_segs):
            if seg_idx >= 25:
                print(f"WARNING: Max 25 segments supported by template format. Skipping SEG {seg_num}")
                break
                
            # Horizontal Step Logic: +5 columns per segment
            col_offset = 3 + (seg_idx * 5)  # SEG 1 -> Col 3 (C)
            spans = cable_spans[seg_num]
            
            # Header Population
            if spans:
                first = spans[0]
                sheet.cell(row=9, column=col_offset).value = first.get('METHOD', 'New Conduit')
                sheet.cell(row=10, column=col_offset).value = first.get('START_ADDRESS', '')
                sheet.cell(row=11, column=col_offset).value = first.get('END_ADDRESS', '')
                sheet.cell(row=12, column=col_offset).value = int(first.get('SIZE', 48))
                
            # Scrub template dummy lines first
            for r in range(15, 200):
                for c in range(col_offset, col_offset + 3):
                    cell = sheet.cell(row=r, column=c)
                    if type(cell).__name__ != 'MergedCell':
                        cell.value = None
            
            # Data Layout Logic
            r_idx = 15
            
            # Formulate handhole storage dynamically using spatial Splitter proximity logic
            hh_store = 50
            if spans:
                bx, by = spans[0].get('X', 0), spans[0].get('Y', 0)
                for sx, sy in splitters:
                    if math.hypot(sx - bx, sy - by) < 15.0: # 15 feet tolerance
                        hh_store = 15
                        break
            
            # Anchor Row (Zero Span, Starting HH Storage)
            sheet.cell(row=r_idx, column=col_offset).value = 0 # Span 0
            sheet.cell(row=r_idx, column=col_offset + 1).value = hh_store
            r_idx += 1
            
            # Sequentially layout the actual geometric traces
            for s in spans:
                try:
                    span_val = int(s.get('SPAN', 0))
                except ValueError:
                    span_val = 0
                    
                try:
                    storage_val = int(s.get('STORAGE', 0))
                except ValueError:
                    storage_val = 0
                    
                cell_span = sheet.cell(row=r_idx, column=col_offset)
                if type(cell_span).__name__ != 'MergedCell':
                    cell_span.value = span_val
                    
                cell_store = sheet.cell(row=r_idx, column=col_offset + 1)
                if type(cell_store).__name__ != 'MergedCell':
                    cell_store.value = storage_val
                
                # Apply Center Alignment for the frame
                for c in range(col_offset, col_offset + 3):
                    cell = sheet.cell(row=r_idx, column=c)
                    if type(cell).__name__ != 'MergedCell':
                        cell.alignment = center_align
                
                r_idx += 1

    def populate_labor_span(self, labor_spans):
        print("Populating Labor Span Sheet dynamically...")
        if 'LABOR SPAN SHEET' not in self.wb.sheetnames:
            print("WARNING: LABOR SPAN SHEET not found in workbook.")
            return
            
        sheet = self.wb['LABOR SPAN SHEET']
        from openpyxl.styles import Alignment
        center_align = Alignment(horizontal='center', vertical='center')
        
        # Sort data sequentially by SP number and then alphabetical ITEM#
        # If SP is missing or completely unparseable to integer, set to 999 to throw it to the bottom
        def parse_sp(s):
            try:
                return int(str(s).strip())
            except ValueError:
                return 999
                
        sorted_labor = sorted(labor_spans, key=lambda x: (parse_sp(x.get('SP', 999)), str(x.get('ITEM#', ''))))
        
        r_idx = 7
        for span in sorted_labor:
            sp_val = str(span.get('SP', '')).strip()
            item_val = str(span.get('ITEM#', '')).strip()
            length_val = str(span.get('LENGTH', '0')).strip()
            cond_sz = str(span.get('COND_SZ', '')).strip()
            cond_qty_str = str(span.get('COND_QTY', '0')).strip()
            drop_flg = str(span.get('DROP_FLG', '')).strip()
            
            # Fiber Counts
            f48 = str(span.get('F48', '')).strip()
            f96 = str(span.get('F96', '')).strip()
            f144 = str(span.get('F144', '')).strip()
            f288 = str(span.get('F288', '')).strip()
            f432 = str(span.get('F432', '')).strip()
            
            # Calculate A: JOB PRINT PAGE #
            col_a = f"SP-{sp_val}" if sp_val else ""
            
            # Calculate B: Construction Note Letter Job Print
            col_b = item_val
            
            # Calculate C: SPAN FOOTAGE
            try:
                col_c = int(length_val)
            except ValueError:
                col_c = 0
                
            # Calculate D: NON-STANDARD CONDUIT SIZE
            col_d = cond_sz if cond_sz in ['2', '4', '2"', '4"'] else ""
            
            # Calculate E: IS THIS A PARALLELING DROP CONDUIT?
            col_e = "yes" if drop_flg and drop_flg != '0' else "no"
            
            # Calculate F: CABLES THIS SPAN
            cables = 0
            for f_val in [f48, f96, f144, f288, f432]:
                if f_val and f_val != '0':
                    cables += 1
            col_f = cables
            
            # Calculate G, H, I: CONDUIT CAPACITIES
            try:
                cond_qty = int(cond_qty_str)
            except ValueError:
                cond_qty = 0
                
            col_g = min(cond_qty, 2)
            col_h = max(0, min(cond_qty - 2, 2))
            col_i = max(0, min(cond_qty - 4, 1))
            
            # Map values to columns A through O (1 to 15)
            # A=1, B=2, C=3, D=4, E=5, F=6, G=7, H=8, I=9, J=10, K=11, L=12, M=13, N=14, O=15
            sheet.cell(row=r_idx, column=1).value = col_a
            sheet.cell(row=r_idx, column=2).value = col_b
            sheet.cell(row=r_idx, column=3).value = col_c
            sheet.cell(row=r_idx, column=4).value = col_d
            sheet.cell(row=r_idx, column=5).value = col_e
            sheet.cell(row=r_idx, column=6).value = col_f
            sheet.cell(row=r_idx, column=7).value = col_g
            sheet.cell(row=r_idx, column=8).value = col_h
            sheet.cell(row=r_idx, column=9).value = col_i
            
            # J through O are left blank as per client decision (pending Aerial automation logic later)
            for c in range(10, 16):
                sheet.cell(row=r_idx, column=c).value = ""
                
            # Apply styling
            for col in range(1, 16):
                cell = sheet.cell(row=r_idx, column=col)
                if type(cell).__name__ != 'MergedCell':
                    cell.alignment = center_align
            
            r_idx += 1

    def save(self):
        print(f"Saving changes to {self.output_path}...")
        self.wb.save(self.output_path)
        print("Done!")

if __name__ == '__main__':
    from pipeline import parse_dxf_data
    extracted = parse_dxf_data("1_3_CX_04.02.2026-Final.dxf")
    
    writer = ExcelWriter("MVP_Output_Workbook.xlsx", "HSNC 01-03 WORKBOOK_04.07.2026.xlsx")
    writer.populate_house_count(extracted['house_count'])
    writer.populate_splices(extracted['splitters_1x8'])
    writer.populate_1x4_splits(extracted['splitters_1x8'])
    writer.populate_cable_sheet(extracted)
    writer.populate_labor_span(extracted.get('labor_spans', []))
    writer.save()
