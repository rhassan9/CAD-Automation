import os
import shutil
import openpyxl
import math
import copy
import re
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

class HouseCountRenderer:
    def populate_house_count(self, house_data):
        print("Populating House Count sheet...")
        if 'House Count' not in self.wb.sheetnames: return
        sheet = self.wb['House Count']
        
        calibri_11 = Font(name='Calibri', size=11)
        center_align = Alignment(horizontal='center', vertical='center')
        
        thin_side = Side(style='thin', color='000000')
        thick_side = Side(style='medium', color='000000')
        standard_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
        thick_bottom_border = Border(top=thin_side, bottom=thick_side, left=thin_side, right=thin_side)
        
        unique_houses = {}
        for data in house_data:
            key = f"{data['ADDRESS']}||{data.get('UNIT', '')}"
            if key not in unique_houses:
                unique_houses[key] = data
                
        sorted_houses = sorted(unique_houses.values(), key=lambda x: x['ADDRESS'])
        
        row = 2
        for idx, data in enumerate(sorted_houses):
            is_12th_row = ((idx + 1) % 12 == 0)
            target_border = thick_bottom_border if is_12th_row else standard_border
            
            sheet.cell(row=row, column=1).value = data['ADDRESS']
            sheet.cell(row=row, column=2).value = "FUQUAY-VARINA"
            sheet.cell(row=row, column=3).value = "NORTH CAROLINA"
            sheet.cell(row=row, column=4).value = "27526"
            sheet.cell(row=row, column=5).value = data.get('UNIT', '')
            sheet.cell(row=row, column=6).value = "RES-SFU"
            sheet.cell(row=row, column=7).value = "BURIED"
            sheet.cell(row=row, column=8).value = "Y"
            sheet.cell(row=row, column=9).value = ""
            sheet.cell(row=row, column=10).value = 1
            sheet.cell(row=row, column=11).value = 0
            
            for col in range(1, 12):
                cell = sheet.cell(row=row, column=col)
                cell.font = calibri_11
                cell.alignment = center_align
                cell.border = target_border
                
            row += 1

class SpliceRenderer:
    def populate_splices(self, splitters_1x8):
        print("Populating 1x8 Splice sheet dynamically...")
        sheet_1x8 = self.wb['CABLE TO 1X8 SPLICING']
        self.extracted_1x8_template_data = {}
        
        center_align = Alignment(horizontal='center', vertical='center')

        sorted_splitters = []
        for b_name, struct in splitters_1x8.items():
            if struct.get('placement_no'):
                try:
                    num = int(struct['placement_no'])
                except ValueError:
                    num = 99999
                sorted_splitters.append((num, b_name, struct))
        sorted_splitters.sort(key=lambda x: x[0])

        for (num, base_name, struct) in sorted_splitters:
            placement_no = struct.get('placement_no', '')
            target_1x8_row = None
            for r in range(5, 500):
                cell_val = str(sheet_1x8.cell(row=r, column=1).value).strip() if sheet_1x8.cell(row=r, column=1).value else None
                if cell_val == str(placement_no):
                    target_1x8_row = r
                    break
                    
            if target_1x8_row:
                color_1_fill = copy.copy(sheet_1x8.cell(row=target_1x8_row, column=2).fill)
                color_2_fill = copy.copy(sheet_1x8.cell(row=target_1x8_row, column=3).fill)
                
                self.extracted_1x8_template_data[placement_no] = {
                    'color_1': color_1_fill,
                    'color_2': color_2_fill,
                    'row': target_1x8_row
                }

                c4 = sheet_1x8.cell(row=target_1x8_row, column=4)
                c4.value = int(placement_no)
                c4.alignment = center_align
                
                c5 = sheet_1x8.cell(row=target_1x8_row, column=5)
                c5.value = struct.get('placement_address', '')
                c5.alignment = center_align
                
                c6 = sheet_1x8.cell(row=target_1x8_row, column=6)
                c6.value = base_name
                c6.alignment = center_align

                for port in range(1, 9):
                    out_1x8_str = "SPARE"
                    if port in struct['ports']:
                        out_1x8_str = f"{base_name}, {port} ({base_name}-{port})"
                    c_port = sheet_1x8.cell(row=target_1x8_row, column=6+port)
                    c_port.value = out_1x8_str
                    c_port.alignment = center_align
            else:
                print(f"WARNING: Placement {placement_no} not found in template Column A!")

    def populate_1x4_splits(self, splitters_1x8):
        print("Populating 1x8 to 1x4 Splice sheet dynamically...")
        if 'SPLICING 1X8 TO 1X4 SPLITS' not in self.wb.sheetnames:
            return
            
        sheet_1x4 = self.wb['SPLICING 1X8 TO 1X4 SPLITS']
        header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        header_font = Font(name='Calibri', size=11, bold=True)
        center_align = Alignment(horizontal='center', vertical='center')
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        sorted_splitters = []
        for b_name, struct in splitters_1x8.items():
            if struct.get('placement_no'):
                try:
                    num = int(struct['placement_no'])
                except ValueError:
                    num = 99999
                sorted_splitters.append((num, b_name, struct))
        sorted_splitters.sort(key=lambda x: x[0])
        
        search_start_row = 1
        max_processed_row = 1
        
        for (num, base_name, struct) in sorted_splitters:
            placement_no = struct.get('placement_no', '')
            anchor_row = None
            for r in range(search_start_row, 5000):
                cell_val = str(sheet_1x4.cell(row=r, column=1).value).strip() if sheet_1x4.cell(row=r, column=1).value else ''
                if cell_val == '1X8 SPLITTER PLACEMENT':
                    anchor_row = r + 1
                    search_start_row = anchor_row + 1 
                    break
                    
            if not anchor_row:
                print(f"Info: Anchor row block for placement {placement_no} not found. Building fallback section...")
                bottom_header = 1
                for r in range(1, 4000):
                    cell_val = str(sheet_1x4.cell(row=r, column=1).value).strip() if sheet_1x4.cell(row=r, column=1).value else ''
                    if cell_val == '1X8 SPLITTER PLACEMENT':
                        bottom_header = r
                        
                fallback_start = max(max_processed_row, bottom_header + 12) + 2 
                header_strip_row = fallback_start
                
                sheet_1x4.merge_cells(start_row=header_strip_row, start_column=1, end_row=header_strip_row, end_column=6)
                h_cell = sheet_1x4.cell(row=header_strip_row, column=1)
                h_cell.value = '1X8 SPLITTER PLACEMENT'
                h_cell.fill = header_fill
                h_cell.font = header_font
                h_cell.alignment = center_align
                
                anchor_row = fallback_start + 1
                
                for c in range(1, 7):
                    sheet_1x4.cell(row=header_strip_row, column=c).border = thin_border
                    
                sub_headers = ['SPLITTER NAME', 'T COUNT', '1X4 PLACEMENT ADDRESS', 'SPLIT', 'SERVICING', 'SERVICING', 'SERVICING', 'SERVICING']
                for c in range(1, 9):
                    sheet_1x4.merge_cells(start_row=anchor_row+1, start_column=c, end_row=anchor_row+2, end_column=c)
                    sub = sheet_1x4.cell(row=anchor_row+1, column=c)
                    sub.value = sub_headers[c-1]
                    sub.font = header_font
                    sub.alignment = center_align
                    sheet_1x4.cell(row=anchor_row+1, column=c).border = thin_border
                    sheet_1x4.cell(row=anchor_row+2, column=c).border = thin_border
                    
                search_start_row = anchor_row + 12
            
            max_processed_row = max(max_processed_row, anchor_row + 10)
            
            p_cell = sheet_1x4.cell(row=anchor_row, column=1)
            p_cell.value = int(placement_no)
            p_cell.alignment = center_align
            
            if hasattr(self, 'extracted_1x8_template_data') and placement_no in self.extracted_1x8_template_data:
                import copy
                color_1 = self.extracted_1x8_template_data[placement_no].get('color_1')
                color_2 = self.extracted_1x8_template_data[placement_no].get('color_2')
                if color_1: sheet_1x4.cell(row=anchor_row, column=3).fill = copy.copy(color_1)
                if color_2: sheet_1x4.cell(row=anchor_row, column=4).fill = copy.copy(color_2)
                    
            addr_cell = sheet_1x4.cell(row=anchor_row, column=6)
            addr_cell.value = struct.get('placement_address', '')
            addr_cell.font = Font(name='Calibri', size=11, bold=True)
            addr_cell.alignment = center_align
            
            t_counts = ["1-4", "5-8", "9-12", "13-16", "17-20", "21-24", "25-28", "29-32"]
            for port in range(1, 9):
                row_offset = anchor_row + 2 + port
                
                for c in range(1, 9):
                    sheet_1x4.cell(row=row_offset, column=c).border = thin_border
                    sheet_1x4.cell(row=row_offset, column=c).alignment = center_align
                
                sheet_1x4.cell(row=row_offset, column=2).value = t_counts[port - 1]
                
                if port in struct['ports']:
                    port_data = struct['ports'][port]
                    raw_name = port_data.get('1x4_name', '')
                    clean_name = raw_name.replace('1X4 SPLITTER ', '').strip()
                    sheet_1x4.cell(row=row_offset, column=1).value = clean_name
                    sheet_1x4.cell(row=row_offset, column=3).value = port_data.get('placement_address_1x4', '')
                    sheet_1x4.cell(row=row_offset, column=4).value = port_data.get('split_reference', '')
                    
                    houses = port_data.get('houses', [])
                    for i in range(4):
                        if i < len(houses) and houses[i]:
                            sheet_1x4.cell(row=row_offset, column=5 + i).value = houses[i]
                            
                    # Clean up: If Name (Col 1) and Address (Col 3) exist, fill any empty cell in Col 4-8 with SPARE
                    v1 = sheet_1x4.cell(row=row_offset, column=1).value
                    v3 = sheet_1x4.cell(row=row_offset, column=3).value
                    if v1 and v3:
                        for check_c in range(4, 9):
                            cell_val = sheet_1x4.cell(row=row_offset, column=check_c).value
                            if not cell_val or str(cell_val).strip() == '':
                                sheet_1x4.cell(row=row_offset, column=check_c).value = 'SPARE'
                                
                else:
                    sheet_1x4.cell(row=row_offset, column=4).value = 'SPARE'
                    sheet_1x4.cell(row=row_offset, column=1).value = None
                    sheet_1x4.cell(row=row_offset, column=3).value = None
                    sheet_1x4.cell(row=row_offset, column=5).value = None
                    sheet_1x4.cell(row=row_offset, column=6).value = None
                    sheet_1x4.cell(row=row_offset, column=7).value = None
                    sheet_1x4.cell(row=row_offset, column=8).value = None

class CableSheetRenderer:
    def populate_cable_sheet(self, data):
        # ---------------------------------------------------------------
        # CABLE SHEET AUTO-POPULATION IS DISABLED (v2.0)
        # Reason: CAD data hygiene (incorrect F48/F96/F144 checkboxes,
        # missing attributes, stacked conduits) makes rigid extraction
        # produce incorrect values. A forensic audit of 1_05_CX_4.2.26.dxf
        # confirmed 1,205 mismatch blocks and 4,995 blind blocks.
        #
        # Replacement: Use the Manual Extraction Assistant script
        # (dfx_debugger.py) which produces Cable_Sheet_Assistant.csv
        # with all raw conduit data organized by segment for manual entry.
        # ---------------------------------------------------------------
        print("INFO: Cable Sheet population is DISABLED. Use Cable_Sheet_Assistant.csv for manual data entry.")

class LaborSpanRenderer:
    def populate_labor_span(self, labor_spans):
        print("Populating Labor Span Sheet dynamically...")
        if 'LABOR SPAN SHEET' not in self.wb.sheetnames:
            print("WARNING: LABOR SPAN SHEET not found in workbook.")
            return
            
        sheet = self.wb['LABOR SPAN SHEET']
        center_align = Alignment(horizontal='center', vertical='center')
        
        def parse_sp(s):
            try:
                return int(str(s).strip())
            except ValueError:
                return 999
                
        sorted_labor = sorted(labor_spans, key=lambda x: (parse_sp(x.get('SP', 999)), str(x.get('ITEM#', ''))))
        
        r_idx = 7
        for span in sorted_labor:
            sp_val = str(span.get('SP', '')).strip()
            
            # Ignore blind blocks (like those on CROSSING PROFILES) that have no Job Print Page
            if not sp_val or sp_val.upper() in ['NONE', '0']:
                continue
                
            item_val = str(span.get('ITEM#', '')).strip()
            length_val = str(span.get('LENGTH', '0')).strip()
            cond_sz = str(span.get('COND_SZ', '')).strip()
            cond_qty_str = str(span.get('COND_QTY', '0')).strip()
            drop_flg = str(span.get('DROP_FLG', '')).strip()
            
            f48 = str(span.get('F48', '')).strip()
            f96 = str(span.get('F96', '')).strip()
            f144 = str(span.get('F144', '')).strip()
            f288 = str(span.get('F288', '')).strip()
            f432 = str(span.get('F432', '')).strip()
            
            # --- STRICT EXTRACTION & REGEX LOGIC ---
            col_a = f"SP-{sp_val}" if sp_val else ""
            col_b = item_val
            
            # Bulletproof Length Extraction (Strips text like "B-F ")
            try:
                col_c = int(length_val)
            except ValueError:
                match = re.search(r'\d+', str(length_val))
                col_c = int(match.group()) if match else 0
                
            # Per client review: Column D (Non-Standard Size), Column E (Paralleling Drop), 
            # and Column F (Cables this span) require manual designer entry. Leaving strictly blank.
            col_d = ""
            col_e = ""
            col_f = ""
            
            try:
                cond_qty = int(cond_qty_str)
            except ValueError:
                cond_qty = 0
                
            col_g = min(cond_qty, 2)
            col_h = max(0, min(cond_qty - 2, 2))
            col_i = max(0, min(cond_qty - 4, 1))
            
            sheet.cell(row=r_idx, column=1).value = col_a
            sheet.cell(row=r_idx, column=2).value = col_b
            sheet.cell(row=r_idx, column=3).value = col_c
            sheet.cell(row=r_idx, column=4).value = col_d
            sheet.cell(row=r_idx, column=5).value = col_e
            sheet.cell(row=r_idx, column=6).value = col_f
            sheet.cell(row=r_idx, column=7).value = col_g
            sheet.cell(row=r_idx, column=8).value = col_h
            sheet.cell(row=r_idx, column=9).value = col_i
            
            for c in range(10, 16):
                sheet.cell(row=r_idx, column=c).value = ""
                
            for col in range(1, 16):
                cell = sheet.cell(row=r_idx, column=col)
                if type(cell).__name__ != 'MergedCell':
                    cell.alignment = center_align
            
            r_idx += 1
            
        # Clear any remaining old data in the template to avoid trailing ghost data
        max_row = sheet.max_row
        while r_idx <= max_row:
            for c in range(1, 16):
                sheet.cell(row=r_idx, column=c).value = ""
            r_idx += 1

class ExcelExporter(HouseCountRenderer, SpliceRenderer, CableSheetRenderer, LaborSpanRenderer):
    def __init__(self, output_path, template_path=None):
        self.output_path = output_path
        if template_path and os.path.exists(template_path):
            shutil.copy(template_path, output_path)
            self.wb = openpyxl.load_workbook(output_path)
            print(f"Loaded existing template: {output_path}")
        else:
            self.wb = openpyxl.Workbook()
            sheets = ['House Count', 'CABLE TO 1X8 SPLICING', 'SPLICING 1X8 TO 1X4 SPLITS']
            for name in sheets:
                self.wb.create_sheet(title=name)
            if 'Sheet' in self.wb.sheetnames:
                self.wb.remove(self.wb['Sheet'])
            print(f"Generated Raw Fallback Template: {output_path}")

    def save(self):
        print(f"Saving changes to {self.output_path}...")
        self.wb.save(self.output_path)
        print("Done!")
